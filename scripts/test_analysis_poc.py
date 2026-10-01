"""Rollback-only service contract checks; load candidate modules before executing."""
import frappe
from unittest.mock import patch
from rpm_worklog import analysis, scope, analysis_export
import csv
import io
from openpyxl import load_workbook

assert frappe.local.site == 'frontend' and frappe.conf.get('rpm_worklog_model_poc')

def denied(fn):
    try:
        fn()
    except (frappe.PermissionError, frappe.ValidationError):
        return
    raise AssertionError('Expected denial')

original_user = frappe.session.user
original_company = frappe.conf.get('rpm_worklog_company')
try:
    frappe.set_user('Administrator')
    scope.install_analysis_role()
    scope.install_analysis_role()
    assert frappe.db.exists('Role',scope.MANAGEMENT_ROLE)
    for table in ('DocPerm','Custom DocPerm'):
        assert not frappe.db.exists(table,{'role':scope.MANAGEMENT_ROLE,'write':1})
    employees = frappe.get_all('Employee', filters={'status':'Active'},
        fields=['name','user_id','company'], limit_page_length=0)
    employees = [e for e in employees if e.user_id and e.user_id != 'Administrator']
    manager, own, inactive, indirect, foreign = employees[:5]
    company = manager.company
    assert company and len(employees) >= 5
    frappe.conf.rpm_worklog_company = company
    for employee in (own, inactive, indirect):
        frappe.db.set_value('Employee', employee.name, {'company':company,'reports_to':manager.name})
    frappe.db.set_value('Employee', inactive.name, 'status', 'Left')
    frappe.db.set_value('Employee', indirect.name, 'reports_to', own.name)
    # SQL-only foreign company fixture avoids provisioning unrelated accounting data.
    frappe.db.set_value('Employee', foreign.name, 'company', 'ANALYSIS OTHER COMPANY')
    prefix = 'analysis-' + frappe.generate_hash(length=10)
    expected = {}
    for i, employee in enumerate((manager, own, inactive, indirect, foreign)):
        name = f'{prefix}-{i}'
        expected[employee.name] = name
        frappe.db.sql('''INSERT INTO `tabRPM Daily Work Log`
            (name,owner,employee,employee_name,department,title,work_date,total_hours,review_state,docstatus)
            VALUES (%s,%s,%s,'Stored Name','Stored Department','Analysis test','2099-01-01',8.4,'Approved',0)''',
            (name, employee.user_id, employee.name))
        for idx, hours in ((1,8.0),(2,0.4)):
            frappe.db.sql('''INSERT INTO `tabRPM Work Log Line`
                (name,parent,parenttype,parentfield,idx,hours,work_item,note,quantity,uom,result,
                 activity_type,record_quantity,item_code,item_name_snapshot,item_stock_uom,
                 item_conversion_factor,work_target,start_time,end_time)
                VALUES (%s,%s,'RPM Daily Work Log','lines',%s,%s,'Recorded work','Note',12,'PCS',
                'Completed','Activity',1,'ITEM','Snapshot','PCS',2,'TARGET','09:00:00','09:10:00')''',
                (f'{name}-{idx}',name,idx,hours))
    filters = dict(from_date='2099-01-01',to_date='2099-01-01')
    def query(kind='Organization', **kwargs):
        return analysis.query(kind, filters, **kwargs)
    roles = ['RPM Work Log Pilot','RPM Work Log Manager Pilot',scope.MANAGEMENT_ROLE]
    frappe.set_user(manager.user_id)
    with patch.object(frappe, 'get_roles', return_value=roles):
        data = query(page_size=1)
        assert data['total_rows'] == 8 and data['total_logs'] == 4
        assert abs(data['reported_hours'] - 33.6) < 1e-8 and data['has_more']
        rows = query(page_size=1000)['rows']
        assert len(rows) == 8 and {r['employee'] for r in rows} == {manager.name,own.name,inactive.name,indirect.name}
        assert set(rows[0]) == set(analysis.COLUMNS)
        assert rows[0]['employee_name']=='Stored Name' and rows[0]['department']=='Stored Department'
        assert rows[0]['hours'] in (8.0,0.4) and rows[0]['item_name_snapshot']=='Snapshot'
        assert rows[0]['record_quantity']==1 and rows[0]['start_time']=='9:00:00'
        for field, value in {'work_item':'Recorded work','note':'Note','quantity':12,
                'uom':'PCS','result':'Completed','activity_type':'Activity','item_code':'ITEM',
                'item_stock_uom':'PCS','item_conversion_factor':2,'work_target':'TARGET',
                'end_time':'9:10:00','current_company':company}.items():
            assert rows[0][field] == value, field
        assert query('Self')['total_rows']==2
        assert query('Team')['total_rows']==2  # active direct only
        multi={**filters,'employee':[own.name,inactive.name,own.name], 'department':['Stored Department','Other Department']}
        result=analysis.query('Organization',multi)
        assert result['total_rows']==4 and result['total_logs']==2
        assert len(analysis.dataset('Organization',multi,complete=True)['rows'])==4
        assert analysis.query('Organization',{**multi,'department':['Other Department']})['total_rows']==0
        assert query()['total_rows']==analysis.query('Organization',{**filters,'employee':[],'department':[]})['total_rows']
        assert analysis.query('Organization',{**filters,'employee':[own.name,foreign.name]})['total_rows']==2
        for value in ([own.name]*101,[None],{'x':1},[[]]):
            denied(lambda: analysis.query('Organization',{**filters,'employee':value}))
        frappe.db.set_value('Employee',own.name,{'employee_number':'TEST-MULTI-001','employee_name':'Multi Name'})
        assert analysis.search_filters('Organization','employee','TEST-MULTI-001')['items'][0]['value']==own.name
        assert analysis.search_filters('Organization','employee','Multi Name')['items'][0]['value']==own.name
        assert analysis.search_filters('Organization','employee',inactive.name)['items'][0]['value']==inactive.name
        assert analysis.search_filters('Organization','employee',foreign.name)['items']==[]
        assert analysis.search_filters('Team','employee',inactive.name)['items']==[]
        assert analysis.search_filters('Organization','department','Stored')['items']==[{'value':'Stored Department','label':'Stored Department'}]
        assert analysis.search_filters('Organization','department',"' OR 1=1 --")['items']==[]
        assert analysis.search_filters('Organization','employee','%')['items']==[]
        denied(lambda: analysis.search_filters('Organization','user',''))
        assert query(offset=8)['rows']==[] and query(offset=8)['total_rows']==8
        for bad in ({**filters,'company':company},{**filters,'user':'Administrator'},
                    {**filters,'review_state':'invalid'}):
            denied(lambda: analysis.query('Organization',bad))
        denied(lambda: query(page_size=0))
        denied(lambda: query(offset=-1))
        denied(lambda: query('Unknown'))
        assert analysis.query('Organization',{**filters,'employee':foreign.name})['total_rows']==0
        assert analysis.query('Organization',{**filters,'department':'Current Department'})['total_rows']==0
        assert analysis.query('Organization',{**filters,'employee':"' OR 1=1 --"})['total_rows']==0
        assert analysis.query('Organization',dict(from_date='2099-01-02',to_date='2099-01-02'))['total_rows']==0
        denied(lambda: analysis.query('Organization',dict(from_date='2099-01-02',to_date='2099-01-01')))
        denied(lambda: analysis.query('Organization',dict(from_date='2099-01-01',to_date='2100-01-02')))
        # Missing parent employee, cancelled parent and malformed child links cannot leak.
        frappe.db.set_value('RPM Daily Work Log',expected[own.name],'docstatus',2)
        assert query()['total_rows']==6
        frappe.db.set_value('RPM Daily Work Log',expected[own.name],'docstatus',0)
        frappe.db.set_value('RPM Work Log Line',expected[own.name]+'-1','parentfield','other')
        assert query()['total_rows']==7
        frappe.db.set_value('RPM Work Log Line',expected[own.name]+'-1','parentfield','lines')
        frappe.db.set_value('Employee',inactive.name,'company',None)
        assert query()['total_rows']==6
        frappe.db.set_value('Employee',inactive.name,'company',company)
        frappe.db.set_value('RPM Work Log Line',expected[own.name]+'-1','note',None)
        detail=analysis.query('Organization',{**filters,'employee':own.name})['rows']
        assert detail[0]['note'] is None
        # Full dataset beyond the old viewer cap, stable pages and complete totals.
        for idx in range(3,305):
            name=expected[own.name]
            frappe.db.sql('''INSERT INTO `tabRPM Work Log Line`
                (name,parent,parenttype,parentfield,idx,hours,result)
                VALUES (%s,%s,'RPM Daily Work Log','lines',%s,1,'Completed')''',
                (f'{name}-{idx}',name,idx))
        pages=[query(offset=offset,page_size=100) for offset in (0,100,200,300)]
        ids=[row['line_id'] for page in pages for row in page['rows']]
        assert len(ids)==len(set(ids))==310
        assert all(page['total_rows']==310 and abs(page['reported_hours']-335.6)<1e-8 for page in pages)
        assert not pages[-1]['has_more']
        complete=analysis.dataset('Organization',filters,complete=True)
        assert len(complete['rows'])==310
        csv_rows=list(csv.reader(io.StringIO(analysis_export.serialize(complete,'CSV').decode('utf-8-sig'))))
        assert len(csv_rows)==311 and len(csv_rows[0])==28
        assert {row[0] for row in csv_rows[1:]}=={expected[e.name] for e in (manager,own,inactive,indirect)}
        workbook=load_workbook(io.BytesIO(analysis_export.serialize(complete,'XLSX')),read_only=True)
        xlsx_rows=list(workbook.active.values)
        assert len(xlsx_rows)==311 and tuple(csv_rows[0])==xlsx_rows[0]
        assert abs(sum(row[analysis.COLUMNS.index('hours')] for row in xlsx_rows[1:])-335.6)<1e-8
        workbook.close()
        literal={key:None for key in analysis.COLUMNS}
        literal['note']='=HYPERLINK("https://example.invalid")'
        literal['hours']=8.4
        sample={'rows':[literal]}
        escaped=list(csv.reader(io.StringIO(analysis_export.serialize(sample,'CSV').decode('utf-8-sig'))))[1]
        assert escaped[analysis.COLUMNS.index('note')].startswith("'=")
        workbook=load_workbook(io.BytesIO(analysis_export.serialize(sample,'XLSX')))
        cell=workbook.active.cell(2,analysis.COLUMNS.index('note')+1)
        assert cell.value==literal['note'] and cell.data_type=='s'
        workbook.close()
        assert 'Organization' in analysis.options()['scopes']
        denied(lambda: analysis_export.download('Organization',filters,'HTML'))
        frappe.conf.rpm_worklog_company = None
        denied(query)
        frappe.conf.rpm_worklog_company = 'MISSING COMPANY'
        denied(query)
        frappe.conf.rpm_worklog_company = company
        frappe.db.set_value('Employee', manager.name, 'company', 'ANALYSIS OTHER COMPANY')
        denied(query)
        frappe.db.set_value('Employee', manager.name, 'company', company)
        frappe.db.set_value('Employee', manager.name, 'status', 'Left')
        denied(query)
        frappe.db.set_value('Employee', manager.name, 'status', 'Active')
        saved_user=own.user_id
        frappe.db.set_value('Employee',own.name,'user_id',manager.user_id)
        denied(query)
        frappe.db.set_value('Employee',own.name,'user_id',saved_user)
    with patch.object(frappe, 'get_roles', return_value=['RPM Work Log Pilot']):
        denied(query)
        denied(lambda: analysis_export.download('Organization',filters))
        assert 'Organization' not in analysis.options()['scopes']
    with patch.object(frappe, 'get_roles', return_value=roles):
        frappe.db.set_value('User', manager.user_id, 'enabled', 0)
        denied(query)
    print('PASS analysis: company isolation, inactive history, self/team, fields, totals, filters, paging, fail-closed')
finally:
    frappe.db.rollback()
    frappe.conf.rpm_worklog_company = original_company
    frappe.set_user(original_user)
